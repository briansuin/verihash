import time
import os
import hashlib
from threading import Thread, Lock
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
import database
import docx
import fitz
import openpyxl
import config_manager

config = config_manager.load_config()
WORKSPACE_DIR = config.get("watchdog_target_dir", "./test_workspace")
DEBOUNCE_WAIT = 60.0

cooling_dict = {}
dict_lock = Lock()

class DebounceHandler(FileSystemEventHandler):
    def _handle_event(self, event):
        if event.is_directory:
            return
            
        # Ignore meaningless directories and Word temporary files
        filename = os.path.basename(event.src_path)
        if any(ignored in event.src_path for ignored in ['.git', '__pycache__']) or \
           filename.startswith('~$') or \
           filename.endswith('.tmp'):
            return
            
        with dict_lock:
            cooling_dict[event.src_path] = time.time()
            print(f"[DEBOUNCE] Intercepted: {event.src_path} | Resets timer.")

    def on_modified(self, event):
        self._handle_event(event)

    def on_created(self, event):
        self._handle_event(event)

def compute_diff(file_path: str) -> str:
    """
    Phase 7.5: Universal Parser Upgrade.
    Extracts pure text from multi-format files, applying character truncation.
    Supported: txt, md, js, py, json, csv, html, docx, pdf, xlsx.
    Falls back to binary SHA-256 with [BINARY FILE MODIFIED] for non-supported extensions.
    """
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()
    MAX_CHARS = 100000
    
    try:
        # 1. Code & Text format
        if ext in ['.txt', '.md', '.js', '.py', '.json', '.csv', '.html']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
                
        # 2. DOCX format
        elif ext == '.docx':
            doc = docx.Document(file_path)
            content = "\n".join([para.text for para in doc.paragraphs])
            
        # 3. PDF format
        elif ext == '.pdf':
            with fitz.open(file_path) as doc:
                content = ""
                for page in doc:
                    content += page.get_text()
                    
        # 4. XLSX format
        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(file_path)
            text_chunks = []
            for sheet in wb.worksheets:
                text_chunks.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join([str(cell) for cell in row if cell is not None])
                    if row_text:
                        text_chunks.append(row_text)
            content = "\n".join(text_chunks)
            
        # 5. Unknown format: Binary Hash Fallback
        else:
            hasher = hashlib.sha256()
            with open(file_path, 'rb') as f:
                while chunk := f.read(8192):
                    hasher.update(chunk)
            return f"[BINARY FILE MODIFIED: Hash={hasher.hexdigest()}]"
            
        # Character Truncation Protection
        if len(content) > MAX_CHARS:
            return content[:MAX_CHARS] + "...[TRUNCATED]"
            
        return content
            
    except Exception as e:
        error_msg = f"[ERROR] Parser Middleware Exception reading {file_path} | Details: {e}"
        print(f"\033[91m{error_msg}\033[0m")
        return error_msg

def daemon_worker():
    print(f"[DAEMON] Worker thread active. Debounce window: {DEBOUNCE_WAIT}s")
    while True:
        time.sleep(1)
        now = time.time()
        
        ready_files = []
        with dict_lock:
            # Check for files stationary > 60s
            for file_path, last_active in list(cooling_dict.items()):
                if now - last_active >= DEBOUNCE_WAIT:
                    ready_files.append(file_path)
                    del cooling_dict[file_path]
        
        for file_path in ready_files:
            # Prevent processing of files deleted during the debounce period
            if not os.path.exists(file_path):
                print(f"\033[93m[DAEMON] File vanished during debounce. Skipping: {file_path}\033[0m")
                continue
                
            print(f"[DAEMON] Stationary > {DEBOUNCE_WAIT}s: {file_path}. Processing...")
            diff_data = compute_diff(file_path)
            database.commit_snapshot(file_path, diff_data)

def main():
    print("[INIT] Bootstrapping VeriHash Phase 1: Physical Infrastructure...")
        
    database.init_db()
    
    # Start worker daemon
    worker_thread = Thread(target=daemon_worker, daemon=True)
    worker_thread.start()

    # Start filesystem observer
    event_handler = DebounceHandler()
    observer = Observer()
    
    # Multi-Directory Surveillance Patch
    target_dirs = [d.strip() for d in WORKSPACE_DIR.split(',')]
    
    for target in target_dirs:
        if not target:
            continue
            
        if not os.path.exists(target):
            try:
                os.makedirs(target)
            except Exception as e:
                print(f"\033[93m[WARNING] Failed to create directory {target}: {e}. Skipping.\033[0m")
                continue
                
        observer.schedule(event_handler, target, recursive=True)
        print(f"[WATCHER] Observer attached to: {target}")
        
    observer.start()
    
    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        print("\n[SHUTDOWN] Terminating Daemon...")
        observer.stop()
    observer.join()

if __name__ == "__main__":
    main()
